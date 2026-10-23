import io
import csv
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from database import get_db_connection

def export_transactions_to_excel(start_date=None, end_date=None, account_id=None, category_id=None, member_id=None, tx_type=None):
    conn = get_db_connection()
    cursor = conn.cursor()

    query = '''
        SELECT 
            t.id,
            t.transaction_date,
            t.type,
            t.amount,
            c.name as category_name,
            a.name as account_name,
            ta.name as to_account_name,
            m.name as member_name,
            t.note
        FROM transactions t
        LEFT JOIN categories c ON t.category_id = c.id
        LEFT JOIN accounts a ON t.account_id = a.id
        LEFT JOIN accounts ta ON t.to_account_id = ta.id
        LEFT JOIN members m ON t.member_id = m.id
        WHERE 1=1
    '''
    params = []
    if start_date:
        query += " AND t.transaction_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND t.transaction_date <= ?"
        params.append(end_date)
    if account_id:
        query += " AND (t.account_id = ? OR t.to_account_id = ?)"
        params.extend([account_id, account_id])
    if category_id:
        query += " AND t.category_id = ?"
        params.append(category_id)
    if member_id:
        query += " AND t.member_id = ?"
        params.append(member_id)
    if tx_type:
        query += " AND t.type = ?"
        params.append(tx_type)

    query += " ORDER BY t.transaction_date DESC, t.id DESC"
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "Sổ Thu Chi Gia Đình"

    # Title styling
    title_font = Font(name="Segoe UI", size=16, bold=True, color="1E3A8A")
    ws.merge_cells("A1:H1")
    ws["A1"] = "BÁO CÁO THU CHI TÀI CHÍNH GIA ĐÌNH"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # Subtitle / Filter info
    ws.merge_cells("A2:H2")
    ws["A2"] = f"Thời gian trích xuất: {start_date or 'Tất cả'} đến {end_date or 'Hiện tại'}"
    ws["A2"].font = Font(name="Segoe UI", size=10, italic=True, color="64748B")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "STT", "Ngày GD", "Loại", "Số tiền (VNĐ)", "Hạng mục / Danh mục", 
        "Tài khoản / Ví", "Thành viên", "Ghi chú"
    ]
    header_fill = PatternFill(start_color="3B82F6", end_color="3B82F6", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    ws.row_dimensions[4].height = 26

    # Data Rows
    type_map = {
        'income': 'Thu nhập',
        'expense': 'Chi tiêu',
        'transfer': 'Chuyển khoản'
    }

    total_income = 0
    total_expense = 0

    current_row = 5
    for idx, r in enumerate(rows, 1):
        amt = r['amount']
        t_type = r['type']
        t_label = type_map.get(t_type, t_type)

        if t_type == 'income':
            total_income += amt
        elif t_type == 'expense':
            total_expense += amt

        # Account display
        acc_str = r['account_name'] or ''
        if t_type == 'transfer' and r['to_account_name']:
            acc_str += f" ➔ {r['to_account_name']}"

        ws.cell(row=current_row, column=1, value=idx).alignment = Alignment(horizontal="center")
        ws.cell(row=current_row, column=2, value=r['transaction_date']).alignment = Alignment(horizontal="center")
        
        type_cell = ws.cell(row=current_row, column=3, value=t_label)
        type_cell.alignment = Alignment(horizontal="center")
        if t_type == 'income':
            type_cell.font = Font(name="Segoe UI", bold=True, color="16A34A")
        elif t_type == 'expense':
            type_cell.font = Font(name="Segoe UI", bold=True, color="DC2626")
        else:
            type_cell.font = Font(name="Segoe UI", bold=True, color="2563EB")

        amt_cell = ws.cell(row=current_row, column=4, value=amt)
        amt_cell.number_format = '#,##0'
        amt_cell.alignment = Alignment(horizontal="right")

        ws.cell(row=current_row, column=5, value=r['category_name'] or 'Chuyển quỹ/Ví')
        ws.cell(row=current_row, column=6, value=acc_str)
        ws.cell(row=current_row, column=7, value=r['member_name'] or 'Chung')
        ws.cell(row=current_row, column=8, value=r['note'] or '')

        for col in range(1, 9):
            ws.cell(row=current_row, column=col).border = thin_border
            ws.cell(row=current_row, column=col).font = Font(name="Segoe UI", size=10)

        current_row += 1

    # Summary section at the bottom
    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    sum_lbl = ws.cell(row=current_row, column=1, value="TỔNG THU NHẬP:")
    sum_lbl.font = Font(name="Segoe UI", bold=True, color="16A34A")
    sum_lbl.alignment = Alignment(horizontal="right")
    inc_sum = ws.cell(row=current_row, column=4, value=total_income)
    inc_sum.number_format = '#,##0 "₫"'
    inc_sum.font = Font(name="Segoe UI", bold=True, color="16A34A")

    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    exp_lbl = ws.cell(row=current_row, column=1, value="TỔNG CHI TIÊU:")
    exp_lbl.font = Font(name="Segoe UI", bold=True, color="DC2626")
    exp_lbl.alignment = Alignment(horizontal="right")
    exp_sum = ws.cell(row=current_row, column=4, value=total_expense)
    exp_sum.number_format = '#,##0 "₫"'
    exp_sum.font = Font(name="Segoe UI", bold=True, color="DC2626")

    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
    bal_lbl = ws.cell(row=current_row, column=1, value="DƯ TÍCH LŨY (THU - CHI):")
    bal_lbl.font = Font(name="Segoe UI", bold=True, color="1D4ED8")
    bal_lbl.alignment = Alignment(horizontal="right")
    bal_sum = ws.cell(row=current_row, column=4, value=total_income - total_expense)
    bal_sum.number_format = '#,##0 "₫"'
    bal_sum.font = Font(name="Segoe UI", bold=True, color="1D4ED8")

    # Column widths
    col_widths = {1: 8, 2: 14, 3: 15, 4: 18, 5: 25, 6: 25, 7: 20, 8: 30}
    for col, width in col_widths.items():
        ws.column_dimensions[chr(64 + col)].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output

def export_transactions_to_csv(start_date=None, end_date=None, account_id=None, category_id=None, member_id=None, tx_type=None):
    conn = get_db_connection()
    cursor = conn.cursor()

    query = '''
        SELECT 
            t.id,
            t.transaction_date,
            t.type,
            t.amount,
            c.name as category_name,
            a.name as account_name,
            ta.name as to_account_name,
            m.name as member_name,
            t.note
        FROM transactions t
        LEFT JOIN categories c ON t.category_id = c.id
        LEFT JOIN accounts a ON t.account_id = a.id
        LEFT JOIN accounts ta ON t.to_account_id = ta.id
        LEFT JOIN members m ON t.member_id = m.id
        WHERE 1=1
    '''
    params = []
    if start_date:
        query += " AND t.transaction_date >= ?"
        params.append(start_date)
    if end_date:
        query += " AND t.transaction_date <= ?"
        params.append(end_date)
    if account_id:
        query += " AND (t.account_id = ? OR t.to_account_id = ?)"
        params.extend([account_id, account_id])
    if category_id:
        query += " AND t.category_id = ?"
        params.append(category_id)
    if member_id:
        query += " AND t.member_id = ?"
        params.append(member_id)
    if tx_type:
        query += " AND t.type = ?"
        params.append(tx_type)

    query += " ORDER BY t.transaction_date DESC, t.id DESC"
    cursor.execute(query, params)
    rows = cursor.fetchall()
    conn.close()

    output = io.StringIO()
    # Write UTF-8 BOM so Excel opens it with right encoding
    output.write('\ufeff')
    writer = csv.writer(output)

    writer.writerow(["Mã GD", "Ngày", "Loại", "Số tiền (VNĐ)", "Hạng mục", "Tài khoản", "Người thực hiện", "Ghi chú"])

    type_map = {'income': 'Thu nhập', 'expense': 'Chi tiêu', 'transfer': 'Chuyển khoản'}

    for r in rows:
        acc_str = r['account_name'] or ''
        if r['type'] == 'transfer' and r['to_account_name']:
            acc_str += f" -> {r['to_account_name']}"

        writer.writerow([
            r['id'],
            r['transaction_date'],
            type_map.get(r['type'], r['type']),
            int(r['amount']),
            r['category_name'] or 'Chuyển quỹ',
            acc_str,
            r['member_name'] or 'Chung',
            r['note'] or ''
        ])

    return io.BytesIO(output.getvalue().encode('utf-8'))
